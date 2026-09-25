import json
import io
import threading
import unittest
from zipfile import ZipFile
from datetime import datetime
from http.client import HTTPConnection
from http.server import ThreadingHTTPServer
from openpyxl import Workbook

from statistical_agents.api import ApplicationHandler, read_excel, validate_theme_config


class ApiSmokeTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.server = ThreadingHTTPServer(("127.0.0.1", 0), ApplicationHandler)
        cls.thread = threading.Thread(target=cls.server.serve_forever, daemon=True)
        cls.thread.start()

    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown()

    def test_sample_and_analyze(self):
        connection = HTTPConnection("127.0.0.1", self.server.server_port)
        connection.request("GET", "/")
        self.assertEqual(connection.getresponse().status, 200)
        connection.request("GET", "/api/sample")
        sample = json.loads(connection.getresponse().read())
        self.assertGreater(len(sample["rows"]), 0)
        body = json.dumps({"question": "Test", "rows": sample["rows"], "columns": ["segment", "puan"]})
        connection.request("POST", "/api/analyze", body, {"Content-Type": "application/json"})
        response = connection.getresponse()
        self.assertEqual(response.status, 200)
        self.assertEqual(len(json.loads(response.read())["results"]), 3)

    def test_invalid_request_is_400(self):
        connection = HTTPConnection("127.0.0.1", self.server.server_port)
        connection.request("POST", "/api/analyze", "{}", {"Content-Type": "application/json"})
        self.assertEqual(connection.getresponse().status, 400)

    def test_report_is_downloadable_docx(self):
        connection = HTTPConnection("127.0.0.1", self.server.server_port)
        body = json.dumps({"question": "Rapor", "rows": [{"yaş": 24, "grup": "A"}, {"yaş": 35, "grup": "B"}],
                           "columns": ["yaş", "grup"], "relationship_pairs": [["yaş", "grup"]]})
        connection.request("POST", "/api/report", body, {"Content-Type": "application/json"})
        response = connection.getresponse()
        content = response.read()
        self.assertEqual(response.status, 200)
        self.assertEqual(content[:2], b"PK")
        self.assertIn("attachment", response.getheader("Content-Disposition"))

    def test_xlsx_upload_returns_rows_and_rejects_xls(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["zaman", "segment"])
        sheet.append([datetime(2026, 8, 23, 18, 42, 4), "A"])
        sheet.append([35, "B"])
        stream = io.BytesIO()
        workbook.save(stream)
        boundary = "----istatistik-test"
        body = (
            f"--{boundary}\r\nContent-Disposition: form-data; name=\"question\"\r\n\r\nExcel test\r\n"
            f"--{boundary}\r\nContent-Disposition: form-data; name=\"columns\"\r\n\r\n[\"zaman\",\"segment\"]\r\n"
            f"--{boundary}\r\nContent-Disposition: form-data; name=\"relationship_pairs\"\r\n\r\n[[\"zaman\",\"segment\"]]\r\n"
            f"--{boundary}\r\nContent-Disposition: form-data; name=\"file\"; filename=\"anket.xlsx\"\r\n"
            "Content-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet\r\n\r\n"
        ).encode() + stream.getvalue() + f"\r\n--{boundary}--\r\n".encode()
        connection = HTTPConnection("127.0.0.1", self.server.server_port)
        connection.request("POST", "/api/upload", body, {"Content-Type": f"multipart/form-data; boundary={boundary}"})
        response = connection.getresponse()
        data = json.loads(response.read())
        self.assertEqual(response.status, 200)
        self.assertEqual(data["rows"][0]["zaman"], "2026-08-23T18:42:04")
        self.assertEqual(data["columns"], ["zaman", "segment"])
        connection.request("POST", "/api/analyze", body, {"Content-Type": f"multipart/form-data; boundary={boundary}"})
        analyze_response = connection.getresponse()
        self.assertEqual(analyze_response.status, 200, analyze_response.read().decode())
        connection.request("POST", "/api/report", body, {"Content-Type": f"multipart/form-data; boundary={boundary}"})
        self.assertEqual(connection.getresponse().status, 200)

        bad_body = body.replace(b"anket.xlsx", b"anket.xls")
        connection.request("POST", "/api/upload", bad_body, {"Content-Type": f"multipart/form-data; boundary={boundary}"})
        self.assertEqual(connection.getresponse().status, 400)

    def test_first_row_is_schema_and_duplicate_blank_headers_are_normalized(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["", "segment", "segment"])
        sheet.append(["first-data-value", "A", 10])
        stream = io.BytesIO()
        workbook.save(stream)
        rows, mappings = read_excel(stream.getvalue(), "headers.xlsx",
                                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertEqual(list(rows[0]), ["Column 1", "segment", "segment (2)"])
        self.assertEqual(rows[0]["Column 1"], "first-data-value")
        self.assertEqual(len(rows), 1)
        self.assertEqual(mappings[0], {"original": "(boş başlık 1)", "normalized": "Column 1"})
        self.assertEqual(mappings[2]["normalized"], "segment (2)")

    def test_sequential_workbooks_replace_headers(self):
        def workbook_bytes(header):
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(header)
            sheet.append([f"value-{header[0]}", "x"])
            stream = io.BytesIO()
            workbook.save(stream)
            return stream.getvalue()

        first_rows, first_map = read_excel(workbook_bytes(["first_header", "shared"]),
                                            "first.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        second_rows, second_map = read_excel(workbook_bytes(["second_header", "other"]),
                                             "second.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        first_headers = list(first_rows[0])
        second_headers = list(second_rows[0])
        self.assertEqual(first_headers, ["first_header", "shared"])
        self.assertEqual(second_headers, ["second_header", "other"])
        self.assertNotIn("first_header", second_headers)
        self.assertEqual([item["normalized"] for item in second_map], second_headers)

    def test_theme_config_override_api_and_sample_metadata(self):
        config = [{"column": "segment", "enabled": True,
                   "themes": [{"name": "A grubu", "keywords": ["A"]}]}]
        body = json.dumps({"question": "Tema", "rows": [{"segment": "A"}, {"segment": "B"}],
                           "columns": ["segment"], "theme_config": config})
        connection = HTTPConnection("127.0.0.1", self.server.server_port)
        connection.request("POST", "/api/analyze", body, {"Content-Type": "application/json"})
        response = connection.getresponse()
        data = json.loads(response.read())
        self.assertEqual(response.status, 200)
        self.assertEqual(data["results"][0]["metrics"]["theme_analysis"][0]["distribution"][0]["theme"], "A grubu")
        connection.request("GET", "/api/sample")
        sample = json.loads(connection.getresponse().read())
        self.assertIn("theme_config", sample)
        self.assertIsInstance(sample["theme_config"], list)

    def test_theme_analysis_is_in_word_report(self):
        body = json.dumps({"question": "Tema raporu", "rows": [{"yorum": "çok iyi"}, {"yorum": "diğer"}],
                           "columns": ["yorum"], "theme_config": [{"column": "yorum", "enabled": True,
                           "themes": [{"name": "Olumlu", "keywords": ["iyi"]}]}]})
        connection = HTTPConnection("127.0.0.1", self.server.server_port)
        connection.request("POST", "/api/report", body, {"Content-Type": "application/json"})
        response = connection.getresponse()
        content = response.read()
        self.assertEqual(response.status, 200)
        with ZipFile(io.BytesIO(content)) as archive:
            document = archive.read("word/document.xml").decode("utf-8")
        self.assertIn("Tema analizi", document)
        self.assertIn("Olumlu", document)
